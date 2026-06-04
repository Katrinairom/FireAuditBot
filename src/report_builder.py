# src/report_builder.py
import os
from docx import Document
from docx.shared import Pt


def generate_report(template_path, output_path, extracted_data):
    """
    Modifies the Word template by updating the footer and filling data tables.
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Word template not found at: {template_path}")

    doc = Document(template_path)
    store_code = str(extracted_data.get("store_code", "UNKNOWN"))

    # --- 1. REPLACE FOOTER '1234' WITH ACTUAL STORE CODE ---
    for section in doc.sections:
        # Check primary footer, first page footer, and even page footers
        for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
            if footer is not None:
                for para in footer.paragraphs:
                    if "1234" in para.text:
                        for run in para.runs:
                            if "1234" in run.text:
                                run.text = run.text.replace("1234", store_code)
                                run.font.name = 'Calibri'
                                run.font.size = Pt(12)

    # --- 2. POPULATE WORD TABLES DYNAMICALLY BY COLUMNS ---
    occ_rows = extracted_data.get("occupancy_table", [])

    for table in doc.tables:
        # Find the Occupancy Table by evaluating header names in row 0
        headers = [cell.text.strip().lower() for cell in table.rows[0].cells]

        if "area occupied" in headers or "occupancy load" in headers:
            # Match columns dynamically based on where the text is found
            area_col_idx = next((i for i, h in enumerate(
                headers) if "area occupied" in h), None)
            load_col_idx = next((i for i, h in enumerate(
                headers) if "occupancy load" in h), None)

            # Drop data down the matching columns starting from row 1
            for idx, data_row in enumerate(occ_rows):
                if (idx + 1) < len(table.rows):
                    row_cells = table.rows[idx + 1].cells

                    if area_col_idx is not None:
                        cell_p = row_cells[area_col_idx].paragraphs[0]
                        cell_p.text = str(
                            data_row["area_occupied"] if data_row["area_occupied"] is not None else "")
                        cell_p.runs[0].font.name = 'Calibri'
                        cell_p.runs[0].font.size = Pt(12)

                    if load_col_idx is not None:
                        cell_p = row_cells[load_col_idx].paragraphs[0]
                        cell_p.text = str(
                            data_row["occupancy_load"] if data_row["occupancy_load"] is not None else "")
                        cell_p.runs[0].font.name = 'Calibri'
                        cell_p.runs[0].font.size = Pt(12)

        # Look for the Fire Load density section to fill the single calculated variable
        for row in table.rows:
            for cell in row.cells:
                if "fire load density" in cell.text.lower() or "<<fire_load_density>>" in cell.text:
                    # Injects calculated fire load density value directly
                    for para in cell.paragraphs:
                        if "<<fire_load_density>>" in para.text:
                            for run in para.runs:
                                if "<<fire_load_density>>" in run.text:
                                    run.text = run.text.replace("<<fire_load_density>>", str(
                                        extracted_data["fire_load_density"]))
                                    run.font.name = 'Calibri'
                                    run.font.size = Pt(12)

    # Save to the output directory
    doc.save(output_path)
    return output_path
